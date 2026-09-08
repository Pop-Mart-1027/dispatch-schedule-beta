import datetime
import json
import re
from pathlib import Path

import openpyxl

SOURCE = Path(r"C:\Users\老涂\.codex\codex-remote-attachments\01a07d00-4e05-75d2-a412-d23f557f274a\2D8884D4-DC33-4C65-A1BA-8A2B86EC88B4\1-2026-調度組班表.xlsx")
OUTPUT = Path(r"E:\Codex\web-app-beta-5-10-1\public\september-schedules.json")


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def extract(workbook, sheet_name, job_col, id_col, name_col, first_day_col, area_col=None):
    sheet = workbook[sheet_name]
    rows = []
    current_group = ""
    for row_number in range(4, sheet.max_row + 1):
        employee_id = clean(sheet.cell(row_number, id_col).value)
        name = clean(sheet.cell(row_number, name_col).value)
        header = employee_id or clean(sheet.cell(row_number, job_col).value)
        if header and not name:
            current_group = heading_label(header)
            continue
        if not employee_id or not name:
            continue
        if current_group == "行事曆":
            continue
        shifts = [clean(sheet.cell(row_number, column).value) for column in range(first_day_col, first_day_col + 30)]
        if not any(shifts):
            continue
        direct_area = clean(sheet.cell(row_number, area_col).value) if area_col else ""
        rows.append({
            "rowId": f"{sheet_name}-{row_number}",
            "employeeId": employee_id,
            "name": name,
            "title": clean(sheet.cell(row_number, job_col).value),
            "group": f"{direct_area}區" if direct_area else current_group,
            "area": primary_area(shifts),
            "shifts": shifts,
        })
    return rows


def heading_label(value):
    """Keep Excel's group order while omitting vehicle and phone details from the UI heading."""
    label = value.split("(")[0].split("（")[0].strip()
    if "區-" in label:
        return label.split("-", 1)[0]
    if label.startswith("RFQ-1025"):
        return "E區"
    return label


AREA_CODES = sorted([
    "A1", "A2", "B1", "B2", "B3", "B4", "C1", "C2", "D1", "D2", "D3", "E1", "E2",
    "F1", "F2", "G1", "G2", "H1", "H2", "I1", "I2", "I3", "J1", "J2", "K1", "K2", "K3", "K4",
    "L1", "L2", "L3", "L4", "M2", "N1", "N2", "N3", "O1", "O2", "O3", "O4", "P1", "P2", "T1", "T2",
    "U", "V", "W1", "W2", "W3", "X1", "X2", "C", "E", "F", "G", "H", "J", "K", "L", "M", "N", "O", "P", "R", "S", "T",
], key=len, reverse=True)


def primary_area(shifts):
    matches = []
    for shift in shifts:
        for area in AREA_CODES:
            if area in shift:
                matches.append(area)
                break
    if not matches:
        return ""
    return max(set(matches), key=matches.count)


workbook = openpyxl.load_workbook(SOURCE, data_only=True)
payload = {
    "month": "2026-09",
    "days": [str(day) for day in range(1, 31)],
    "morning": extract(workbook, "9月日班", 2, 3, 4, 5, area_col=1),
    "night": extract(workbook, "9月夜班", 1, 2, 3, 4),
}
OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(f"morning={len(payload['morning'])} night={len(payload['night'])} personal={sum(row['employeeId'] == '96504' for row in payload['night'])}")
